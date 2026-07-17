"""Invariant evaluator tests (issue #38).

Covers the acceptance criteria this module exists to satisfy: a stale
cached workbook must not pass solely because formula cells contain no
visible errors, and failed refresh / zero-row load / stale as-of timestamp
/ missing model table / sentinel mismatch must each be individually
detectable in fixtures -- with every declared assertion producing its own
invariant_result entry even when earlier ones already failed.
"""

from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table
from wb_fixtures import save_workbook

from control_plane.errors import ContractError
from control_plane.invariant_evaluator import evaluate_contract

AS_OF_VALUE = datetime(2026, 7, 17, 8, 0, 0)


def _populate_default(wb) -> None:
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "As Of"
    summary["B1"] = AS_OF_VALUE
    summary["A2"] = "Status"
    summary["B2"] = "OK"

    data = wb.create_sheet("Data")
    data.append(["Region", "Revenue"])
    for i in range(5):
        data.append([f"Region{i}", i * 100])

    table = Table(displayName="DataTable", ref="A1:B6")
    data.add_table(table)

    wb.defined_names["SalesRegion"] = DefinedName("SalesRegion", attr_text="Data!$A$1")


PASSING_CONTRACT = {
    "required_sheets": ["Summary", "Data"],
    "required_defined_names": ["SalesRegion"],
    "required_tables": ["DataTable"],
    "min_row_counts": {"Data": 5, "DataTable": 5},
    "sentinel_cells": [{"sheet": "Summary", "cell": "B2", "expected": "OK"}],
    "freshness": {"sheet": "Summary", "cell": "B1", "max_age_hours": 24},
    "prohibit_visible_errors": True,
}


def _by_name(results: list) -> dict:
    return {r["name"]: r for r in results}


def test_passing_contract_reports_every_declared_invariant_as_passed(tmp_path) -> None:
    path = save_workbook(tmp_path / "wb.xlsx", _populate_default)
    now = datetime(2026, 7, 17, 10, 0, 0)  # 2 hours after AS_OF_VALUE

    results = evaluate_contract(path, PASSING_CONTRACT, now=now)

    assert results
    assert all(r["passed"] for r in results)
    by_name = _by_name(results)
    assert by_name["required_sheet:Summary"]["passed"] is True
    assert by_name["required_sheet:Data"]["passed"] is True
    assert by_name["required_defined_name:SalesRegion"]["passed"] is True
    assert by_name["required_table:DataTable"]["passed"] is True
    assert by_name["min_row_count:Data"]["passed"] is True
    assert by_name["min_row_count:DataTable"]["passed"] is True
    assert by_name["sentinel_cell:Summary!B2"]["passed"] is True
    assert by_name["freshness:Summary!B1"]["passed"] is True
    assert by_name["prohibit_visible_errors"]["passed"] is True
    # Every invariant_result matches #34's exact shape: no stray keys.
    for result in results:
        assert set(result.keys()) <= {"name", "passed", "message"}


def test_missing_required_sheet_is_a_failed_invariant_not_an_exception(tmp_path) -> None:
    path = save_workbook(tmp_path / "wb.xlsx", _populate_default)
    contract = {"required_sheets": ["Summary", "DoesNotExist"]}

    results = evaluate_contract(path, contract)

    by_name = _by_name(results)
    assert by_name["required_sheet:Summary"]["passed"] is True
    assert by_name["required_sheet:DoesNotExist"]["passed"] is False
    assert "DoesNotExist" in by_name["required_sheet:DoesNotExist"]["message"]


def test_row_count_violation_is_detected(tmp_path) -> None:
    path = save_workbook(tmp_path / "wb.xlsx", _populate_default)
    contract = {"min_row_counts": {"Data": 100}, "prohibit_visible_errors": False}

    results = evaluate_contract(path, contract)

    assert len(results) == 1
    assert results[0]["name"] == "min_row_count:Data"
    assert results[0]["passed"] is False
    assert "100" in results[0]["message"]


def test_min_row_count_against_an_unknown_name_fails_with_a_clear_message(tmp_path) -> None:
    path = save_workbook(tmp_path / "wb.xlsx", _populate_default)
    contract = {"min_row_counts": {"NoSuchSheetOrTable": 1}, "prohibit_visible_errors": False}

    results = evaluate_contract(path, contract)

    assert results[0]["passed"] is False
    assert "No sheet or table named" in results[0]["message"]


def test_sentinel_cell_mismatch_is_detected(tmp_path) -> None:
    path = save_workbook(tmp_path / "wb.xlsx", _populate_default)
    contract = {
        "sentinel_cells": [{"sheet": "Summary", "cell": "B2", "expected": "FAIL"}],
        "prohibit_visible_errors": False,
    }

    results = evaluate_contract(path, contract)

    assert len(results) == 1
    assert results[0]["name"] == "sentinel_cell:Summary!B2"
    assert results[0]["passed"] is False
    assert "FAIL" in results[0]["message"]
    assert "OK" in results[0]["message"]


def test_stale_freshness_timestamp_fails(tmp_path) -> None:
    path = save_workbook(tmp_path / "wb.xlsx", _populate_default)  # B1 = AS_OF_VALUE, 08:00
    contract = {
        "freshness": {"sheet": "Summary", "cell": "B1", "max_age_hours": 1},
        "prohibit_visible_errors": False,
    }
    now = datetime(2026, 7, 17, 12, 0, 0)  # 4 hours later: stale against a 1-hour window

    results = evaluate_contract(path, contract, now=now)

    assert len(results) == 1
    assert results[0]["name"] == "freshness:Summary!B1"
    assert results[0]["passed"] is False
    assert "hour" in results[0]["message"]


def test_freshness_cell_that_is_not_a_date_fails_clearly(tmp_path) -> None:
    path = save_workbook(tmp_path / "wb.xlsx", _populate_default)  # A1 = "As Of" (a string)
    contract = {
        "freshness": {"sheet": "Summary", "cell": "A1", "max_age_hours": 24},
        "prohibit_visible_errors": False,
    }

    results = evaluate_contract(path, contract)

    assert results[0]["passed"] is False
    assert "datetime/date" in results[0]["message"]


def test_workbook_with_a_prohibited_visible_error_value_is_detected(tmp_path) -> None:
    def populate(wb) -> None:
        _populate_default(wb)
        wb["Summary"]["C1"] = "#REF!"

    path = save_workbook(tmp_path / "wb.xlsx", populate)

    results = evaluate_contract(path, {"prohibit_visible_errors": True})

    assert len(results) == 1
    assert results[0]["name"] == "prohibit_visible_errors"
    assert results[0]["passed"] is False
    assert "#REF!" in results[0]["message"]
    assert "Summary!C1" in results[0]["message"]


def test_prohibit_visible_errors_defaults_to_true_when_omitted(tmp_path) -> None:
    def populate(wb) -> None:
        _populate_default(wb)
        wb["Summary"]["C1"] = "#DIV/0!"

    path = save_workbook(tmp_path / "wb.xlsx", populate)

    results = evaluate_contract(path, {})  # prohibit_visible_errors not declared

    assert len(results) == 1
    assert results[0]["passed"] is False


def test_prohibit_visible_errors_can_be_disabled(tmp_path) -> None:
    def populate(wb) -> None:
        _populate_default(wb)
        wb["Summary"]["C1"] = "#REF!"

    path = save_workbook(tmp_path / "wb.xlsx", populate)

    results = evaluate_contract(path, {"prohibit_visible_errors": False})

    assert results == []


def test_stale_cached_workbook_with_no_visible_errors_still_fails_declared_invariants(
    tmp_path,
) -> None:
    # The acceptance criterion this module exists to satisfy: a workbook can
    # have zero visible formula errors and still be stale/wrong by the
    # contract's own declared invariants (zero-row load here).
    def populate(wb) -> None:
        summary = wb.active
        summary.title = "Summary"
        summary["B1"] = AS_OF_VALUE  # looks fresh
        empty_data = wb.create_sheet("Data")  # zero rows: a failed refresh/load
        assert empty_data  # sheet exists, but has no data

    path = save_workbook(tmp_path / "wb.xlsx", populate)
    contract = {
        "min_row_counts": {"Data": 1},
        "prohibit_visible_errors": True,  # no visible errors anywhere
    }

    results = evaluate_contract(path, contract, now=datetime(2026, 7, 17, 8, 30, 0))

    by_name = _by_name(results)
    assert by_name["prohibit_visible_errors"]["passed"] is True  # no visible errors
    assert by_name["min_row_count:Data"]["passed"] is False  # but the load was empty


def test_expected_calculation_mode_reports_not_checkable_when_openpyxl_cannot_read_it(
    tmp_path,
) -> None:
    path = save_workbook(tmp_path / "wb.xlsx", _populate_default)  # calcMode never set

    results = evaluate_contract(
        path, {"expected_calculation_mode": "manual", "prohibit_visible_errors": False}
    )

    assert len(results) == 1
    assert results[0]["name"] == "expected_calculation_mode"
    assert results[0]["passed"] is False  # fail-closed, not a silent pass
    assert "not_checkable" in results[0]["message"]


def test_expected_calculation_mode_matches_when_checkable(tmp_path) -> None:
    def populate(wb) -> None:
        _populate_default(wb)
        wb.calculation.calcMode = "manual"

    path = save_workbook(tmp_path / "wb.xlsx", populate)

    matching = evaluate_contract(
        path, {"expected_calculation_mode": "manual", "prohibit_visible_errors": False}
    )
    mismatching = evaluate_contract(
        path, {"expected_calculation_mode": "automatic", "prohibit_visible_errors": False}
    )

    assert matching == [{"name": "expected_calculation_mode", "passed": True}]
    assert mismatching[0]["passed"] is False


def test_every_assertion_gets_its_own_entry_even_after_an_earlier_failure(tmp_path) -> None:
    path = save_workbook(tmp_path / "wb.xlsx", _populate_default)
    contract = {
        "required_sheets": ["DoesNotExist"],
        "min_row_counts": {"Data": 100},
        "sentinel_cells": [{"sheet": "Summary", "cell": "B2", "expected": "WRONG"}],
        "prohibit_visible_errors": False,
    }

    results = evaluate_contract(path, contract)

    assert len(results) == 3
    assert all(r["passed"] is False for r in results)


def test_malformed_contract_raises_schema_invalid_not_an_invariant_failure(tmp_path) -> None:
    path = save_workbook(tmp_path / "wb.xlsx", _populate_default)

    with pytest.raises(ContractError) as excinfo:
        evaluate_contract(path, {"required_sheets": "not-a-list"})

    assert excinfo.value.code == "SCHEMA_INVALID"


def test_missing_workbook_raises_workbook_unreadable(tmp_path) -> None:
    with pytest.raises(ContractError) as excinfo:
        evaluate_contract(tmp_path / "does_not_exist.xlsx", {})

    assert excinfo.value.code == "WORKBOOK_UNREADABLE"
